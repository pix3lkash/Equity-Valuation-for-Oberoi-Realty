"""
=========================================================
REPORT GENERATION
Project : Oberoi Realty Stock Valuation

Creates:
    Final_Report.xlsx

=========================================================
"""

from wsgiref import headers

from wsgiref import headers

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image

from config import *


# =========================================================
# Formatting
# =========================================================

TITLE_FILL = PatternFill(
    fill_type="solid",
    start_color="1F4E78",
    end_color="1F4E78"
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    start_color="D9EAD3",
    end_color="D9EAD3"
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    start_color="BDD7EE",
    end_color="BDD7EE"
)

THIN_BORDER = Border(

    left=Side(style="thin"),

    right=Side(style="thin"),

    top=Side(style="thin"),

    bottom=Side(style="thin")

)


# =========================================================
# Title
# =========================================================

def heading(ws, text, row):

    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=6
    )

    cell = ws.cell(row=row, column=1)

    cell.value = text

    cell.font = Font(
        bold=True,
        color="FFFFFF",
        size=16
    )

    cell.fill = TITLE_FILL

    cell.alignment = Alignment(
        horizontal="center"
    )

    return row + 2


# =========================================================
# Section
# =========================================================

def section(ws, text, row):

    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=6
    )

    cell = ws.cell(row=row, column=1)

    cell.value = text

    cell.font = Font(
        bold=True,
        size=13
    )

    cell.fill = SECTION_FILL

    return row + 1


# =========================================================
# Write DataFrame
# =========================================================

def write_df(ws, df, start_row):

    rows = dataframe_to_rows(
        df,
        index=False,
        header=True
    )

    for r_idx, row_data in enumerate(rows, start=start_row):

        for c_idx, value in enumerate(row_data, start=1):

            cell = ws.cell(
                row=r_idx,
                column=c_idx,
                value=value
            )

            cell.border = THIN_BORDER

            if r_idx == start_row:

                cell.font = Font(bold=True)

                cell.fill = HEADER_FILL

    return start_row + len(df) + 3


# =========================================================
# Insert Chart
# =========================================================

def insert_chart(ws, image_path, cell):

    try:

        img = Image(str(image_path))

        img.width = 700

        img.height = 350

        ws.add_image(img, cell)

    except Exception:

        print(f"Chart missing : {image_path}")


# =========================================================
# Auto Fit
# =========================================================

def autofit(ws):

    for col in range(1, ws.max_column + 1):

        letter = get_column_letter(col)

        maximum = 0

        for row in range(1, ws.max_row + 1):

            cell = ws.cell(row=row, column=col)

            if cell.__class__.__name__ == "MergedCell":

                continue

            if cell.value is not None:

                maximum = max(
                    maximum,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            letter
        ].width = min(maximum + 3, 35)

# =========================================================
# Main Report
# =========================================================

def create_report():

    print("\nGenerating Final Report...\n")

    wb = Workbook()

    ws = wb.active

    ws.title = "Equity Report"

    row = 1

    # =====================================================
    # Title
    # =====================================================

    row = heading(
        ws,
        "OBEROI REALTY LTD - EQUITY VALUATION REPORT",
        row
    )

    # =====================================================
    # Load All Data
    # =====================================================

    ratios = pd.read_csv(
        OUTPUT / "Financial_Ratios.csv"
    )

    financial = pd.read_csv(
        PROCESSED_DATA / "Financial_Summary.csv"
    )

    technical = pd.read_csv(
        OUTPUT / "Technical_Indicators.csv"
    )

    dcf = pd.read_csv(
        OUTPUT / "DCF_Valuation.csv"
    )

    ddm = pd.read_csv(
        OUTPUT / "DDM_Valuation.csv"
    )

    capm = pd.read_csv(
        OUTPUT / "CAPM.csv"
    )

    prediction = pd.read_csv(
        OUTPUT / "Price_Prediction.csv"
    )

    recommendation = pd.read_csv(
        OUTPUT / "Recommendation.csv"
    )

    # =====================================================
    # Executive Summary
    # =====================================================

    row = section(
        ws,
        "Executive Summary",
        row
    )

    market_price = float(
        technical["Close"].iloc[-1]
    )

    dcf_value = float(

        dcf.loc[
            dcf["Metric"] ==
            "Intrinsic Value Per Share",
            "Value"
        ].values[0]

    )

    ddm_value = float(

        ddm.loc[
            ddm["Metric"] ==
            "Intrinsic Value",
            "Value"
        ].values[0]

    )

    pe = float(

        ratios.loc[
            ratios["Ratio"] == "P/E",
            "Value"
        ].values[0]

    )

    pb = float(

        ratios.loc[
            ratios["Ratio"] == "P/B",
            "Value"
        ].values[0]

    )

    roe = float(

        ratios.loc[
            ratios["Ratio"] == "ROE",
            "Value"
        ].values[0]

    )

    overall = recommendation.loc[
        recommendation["Method"] == "OVERALL",
        "Signal"
    ].iloc[0]

    summary = pd.DataFrame({

        "Metric": [

            "Current Market Price",

            "DCF Intrinsic Value",

            "DDM Intrinsic Value",

            "P/E Ratio",

            "P/B Ratio",

            "ROE",

            "Overall Recommendation"

        ],

        "Value": [

            round(market_price,2),

            round(dcf_value,2),

            round(ddm_value,2),

            round(pe,2),

            round(pb,2),

            f"{roe*100:.2f}%",

            overall

        ]

    })

    row = write_df(
        ws,
        summary,
        row
    )

    # ===================================================
    # Price Prediction
    # ===================================================

    prediction = pd.read_csv(
        OUTPUT / "Price_Prediction.csv"
    )

    row = section(ws, "Price Prediction", row)

    row = write_df(ws, prediction, row)

    insert_chart(
        ws,
        CHARTS / "predicted_prices.png",
        f"A{row}"
    )

    row += 22

    # ===================================================
    # Valuation Comparison
    # ===================================================

    row = section(ws, "Valuation Comparison - Majority View", row)

    rec = recommendation.set_index("Method")

    headers = [
    "Method",
    "Actual",
    "Benchmark",
    "Signal",
    "Reason"
]

    for col, header in enumerate(headers, start=1):

        cell = ws.cell(row=row, column=col)

        cell.value = header

        cell.font = Font(bold=True)

        cell.fill = HEADER_FILL

        cell.border = THIN_BORDER

    row += 1

    methods = [
    "DCF",
    "DDM",
    "P/E",
    "P/B",
    "P/S",
    "CAPM",
    "Technical"
]

    for method in methods:

        ws.cell(row=row, column=1).value = method

        ws.cell(row=row, column=2).value = rec.loc[method, "Actual"]

        ws.cell(row=row, column=3).value = rec.loc[method, "Benchmark"]

        signal = rec.loc[method, "Signal"]

        signal_cell = ws.cell(row=row, column=4)

        signal_cell.value = signal

        if signal == "BUY":
            signal_cell.font = Font(color="008000", bold=True)

        elif signal == "SELL":
            signal_cell.font = Font(color="FF0000", bold=True)

        else:
            signal_cell.font = Font(color="E69138", bold=True)

        ws.cell(row=row, column=5).value = rec.loc[method, "Reason"]

        for c in range(1, 6):
            ws.cell(row=row, column=c).border = THIN_BORDER

        row += 1

    row += 1

    insert_chart(
    ws,
    CHARTS / "valuation_comparison.png",
    f"A{row}"
)

    row += 22

    # ===================================================
    # Final Recommendation
    # ===================================================

    row = section(ws, "Majority Vote Recommendation", row)

    comparison = recommendation[
        recommendation["Method"] != "OVERALL"
]

    buy_count = (comparison["Signal"] == "BUY").sum()
    hold_count = (comparison["Signal"] == "HOLD").sum()
    sell_count = (comparison["Signal"] == "SELL").sum()

    overall = recommendation.loc[
        recommendation["Method"] == "OVERALL",
        "Signal"
].iloc[0]

    summary = [
    ("BUY Signals", buy_count),
    ("HOLD Signals", hold_count),
    ("SELL Signals", sell_count),
    ("Overall Recommendation", overall)
]

    for label, value in summary:

        ws.cell(row=row, column=1).value = label

        cell = ws.cell(row=row, column=2)

        cell.value = value

        if value == "BUY":
            cell.font = Font(color="008000", bold=True, size=14)

        elif value == "SELL":
            cell.font = Font(color="FF0000", bold=True, size=14)

        elif value == "HOLD":
            cell.font = Font(color="E69138", bold=True, size=14)

        row += 1

    row += 2

    if overall == "BUY":

        message = (
        "The majority of valuation methods indicate that "
        "Oberoi Realty is undervalued. The overall "
        "recommendation is BUY."
    )

    elif overall == "SELL":

        message = (
        "The majority of valuation methods indicate that "
        "Oberoi Realty is overvalued. The overall "
        "recommendation is SELL."
    )

    else:

        message = (
        "The valuation methods provide mixed signals. "
        "The majority vote suggests a HOLD recommendation."
    )

    ws.merge_cells(
    start_row=row,
    start_column=1,
    end_row=row + 3,
    end_column=5
)

    cell = ws.cell(row=row, column=1)

    cell.value = message

    cell.alignment = Alignment(
    wrap_text=True,
    vertical="top"
)

    cell.font = Font(
    italic=True,
    size=11
)

    row += 5

    # ===================================================
    # Auto-fit Columns
    # ===================================================

    for col in range(1, ws.max_column + 1):

        letter = get_column_letter(col)

        max_length = 0

        for row_num in range(1, ws.max_row + 1):

            cell = ws.cell(row=row_num, column=col)

            if cell.__class__.__name__ == "MergedCell":
                continue

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[letter].width = min(
            max_length + 3,
            35
        )

    # ===================================================
    # Save Workbook
    # ===================================================

    output = REPORT / "Final_Report.xlsx"

    wb.save(output)

    print("\n✓ Final Report Created")
    print(output)


if __name__ == "__main__":
    create_report()

